"""Medical MCQ Telegram bot (v2).

What changed vs v1:
  * Fully async: Gemini calls no longer block the event loop (one user can't freeze everyone).
  * Job queue: limited number of heavy jobs at once, 1 active job per user, queue position shown.
  * Big files are split into chunks; chunks are generated in parallel and sent as they finish.
  * Slides: pages that contain real figures/diagrams (or are scanned) are rendered and sent to
    Gemini as images; logos/decoration are ignored. Images are only used as a source of facts.
  * Modes: Auto / MCQ / Clinical cases (case vignette sent as a message, then the quiz poll).
  * LLM pool: several API keys x several models with per-(key,model) cooldowns, automatic retry
    with backoff, and fallback when a model/key hits quota or is overloaded.
  * State lives in SQLite (DB_PATH) instead of JSON files.
"""
import asyncio
import json
import logging
import os
import random
import re
import sqlite3
import time
import uuid
from dataclasses import dataclass, field

import pymupdf as fitz  # PyMuPDF
from dotenv import load_dotenv
from google import genai
from google.genai import types
from pydantic import BaseModel
from telegram import ReplyKeyboardMarkup, Update
from telegram.error import BadRequest, NetworkError, RetryAfter, TimedOut
from telegram.ext import (
    Application, CommandHandler, ContextTypes, MessageHandler,
    PollAnswerHandler, PollHandler, filters,
)

load_dotenv()

# ───────────────────────── config ─────────────────────────
TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN", "")
_keys_raw = os.getenv("GEMINI_API_KEYS") or os.getenv("GEMINI_API_KEY") or ""
GEMINI_KEYS = [k.strip() for k in _keys_raw.split(",") if k.strip()]
GEMINI_MODELS = [
    m.strip() for m in os.getenv(
        "GEMINI_MODELS",
        "gemini-3.1-flash-lite,gemini-3.5-flash-lite,gemini-3-flash-preview",
    ).split(",") if m.strip()
]
LLM_CONCURRENCY = int(os.getenv("LLM_CONCURRENCY", "4"))
JOB_WORKERS = int(os.getenv("JOB_WORKERS", "2"))
MAX_QUEUE = int(os.getenv("MAX_QUEUE", "25"))
MAX_PAGES = int(os.getenv("MAX_PAGES", "150"))
DB_PATH = os.getenv("DB_PATH", "bot_data.db")
LLM_CALL_TIMEOUT = int(os.getenv("LLM_CALL_TIMEOUT", "150"))
LLM_TOTAL_TIMEOUT = int(os.getenv("LLM_TOTAL_TIMEOUT", "300"))
TG_FILE_LIMIT = 20 * 1024 * 1024  # Bot API download limit

MAX_PAGES_PER_CHUNK = 8
MAX_IMAGES_PER_CHUNK = 6
MAX_CHARS_PER_CHUNK = 9000

log = logging.getLogger("mcqbot")

WAITING_BUTTONS = {"🤖 Auto": "auto", "📝 MCQ": "mcq", "🏥 Cases": "cases"}
MODE_HELP = {
    "auto": "🤖 Auto – clinical topics become clinical cases, basic-science topics stay direct MCQs",
    "mcq": "📝 MCQ – direct multiple-choice questions",
    "cases": "🏥 Cases – every question is a clinical case scenario",
}


# ───────────────────────── database ─────────────────────────
db = sqlite3.connect(DB_PATH, check_same_thread=False)
db.row_factory = sqlite3.Row
db.executescript(
    """
    PRAGMA journal_mode=WAL;
    CREATE TABLE IF NOT EXISTS settings(user_id INTEGER PRIMARY KEY, mode TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS quizzes(
        quiz_id TEXT PRIMARY KEY, user_id INTEGER, chat_id INTEGER,
        gen_done INTEGER DEFAULT 0, finished INTEGER DEFAULT 0, created REAL);
    CREATE TABLE IF NOT EXISTS polls(
        poll_id TEXT PRIMARY KEY, quiz_id TEXT, user_id INTEGER, chat_id INTEGER,
        correct_idx INTEGER, answered INTEGER DEFAULT 0, right INTEGER DEFAULT 0);
    CREATE INDEX IF NOT EXISTS polls_quiz ON polls(quiz_id);
    """
)
# housekeeping: nothing is generating right after a restart; drop old data
db.execute("UPDATE quizzes SET gen_done=1 WHERE gen_done=0")
_cut = time.time() - 14 * 86400
db.execute("DELETE FROM polls WHERE quiz_id IN (SELECT quiz_id FROM quizzes WHERE created<?)", (_cut,))
db.execute("DELETE FROM quizzes WHERE created<?", (_cut,))
db.commit()


def get_mode(user_id: int) -> str:
    r = db.execute("SELECT mode FROM settings WHERE user_id=?", (user_id,)).fetchone()
    return r["mode"] if r else "auto"


def set_mode(user_id: int, mode: str) -> None:
    db.execute("INSERT INTO settings(user_id, mode) VALUES(?,?) "
               "ON CONFLICT(user_id) DO UPDATE SET mode=excluded.mode", (user_id, mode))
    db.commit()


# ───────────────────────── question model ─────────────────────────
class QItem(BaseModel):
    case: str            # clinical case text, or "" for a direct question
    question: str
    correct_answer: str  # the ONE correct option text
    explanation: str     # why correct_answer is right
    wrong_answers: list[str]  # exactly 3 plausible wrong options


class Verdict(BaseModel):
    id: int
    answer: str          # exact text of the option the verifier picked


_LABEL_RE = re.compile(r"^\s*[A-Ea-e][\)\.\-:]\s+")
_ABOVE_RE = re.compile(r"\babove\b|\b(all|none) of (the|these)\b|^both\b", re.I)


def clip(s: str, n: int) -> str:
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def normalize_question(q: QItem) -> dict | None:
    """Build the option list ourselves: no answer index can ever point at the wrong option."""
    right = _LABEL_RE.sub("", q.correct_answer or "").strip()
    wrong = [_LABEL_RE.sub("", w).strip() for w in q.wrong_answers if w and w.strip()]
    opts = [right] + wrong
    if not right or not (2 <= len(opts) <= 6) or not q.question.strip():
        return None
    if len({o.lower() for o in opts}) != len(opts) or any(_ABOVE_RE.search(o) for o in opts):
        return None
    random.shuffle(opts)    # models love putting the answer first
    opts = [clip(o, 100) for o in opts]
    right = clip(right, 100)
    return {
        "case": q.case.strip(),
        "question": " ".join(q.question.split()),
        "options": opts,
        "correct": opts.index(right),
        "explanation": clip(q.explanation, 200),
    }


# ───────────────────────── LLM pool ─────────────────────────
class Busy(Exception):
    pass


class LLMPool:
    """Round-robins API keys, falls back across models, honours quota cooldowns."""

    def __init__(self, keys: list[str], models: list[str], concurrency: int):
        if not keys:
            raise SystemExit("No Gemini key. Set GEMINI_API_KEYS (comma separated) in the environment.")
        self.clients = [genai.Client(api_key=k) for k in keys]
        self.models = models
        self.sem = asyncio.Semaphore(concurrency)
        self.cool: dict[tuple[int, str], float] = {}
        self.dead_models: set[str] = set()
        self.rr = 0

    def _available(self) -> list[tuple[int, str]]:
        now = time.monotonic()
        out = []
        for model in self.models:                      # model priority first …
            if model in self.dead_models:
                continue
            for off in range(len(self.clients)):       # … then rotate keys
                ki = (self.rr + off) % len(self.clients)
                if self.cool.get((ki, model), 0) <= now:
                    out.append((ki, model))
        return out

    def _cool(self, ki: int | None, model: str, seconds: float) -> None:
        until = time.monotonic() + seconds
        keys = range(len(self.clients)) if ki is None else [ki]
        for k in keys:
            self.cool[(k, model)] = max(self.cool.get((k, model), 0), until)

    @staticmethod
    def _parse(text: str) -> list[dict]:
        text = (text or "").strip()
        if text.startswith("```"):
            text = text.strip("`")
            if text.lower().startswith("json"):
                text = text[4:]
        data = json.loads(text)
        if isinstance(data, dict):
            data = data.get("questions", [data])
        if not isinstance(data, list):
            raise ValueError("not a list")
        return data

    async def generate(self, contents: list, schema=QItem, temperature: float = 0.4) -> list:
        cfg = types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=list[schema],
            temperature=temperature,
            max_output_tokens=16000,
        )
        deadline = time.monotonic() + LLM_TOTAL_TIMEOUT
        parse_fails = 0
        last_err = None
        while time.monotonic() < deadline:
            cands = self._available()
            if not cands:
                live = [t for (k, m), t in self.cool.items() if m not in self.dead_models]
                if not live or len(self.dead_models) >= len(self.models):
                    break
                wait = max(0.5, min(live) - time.monotonic())
                if time.monotonic() + wait > deadline:
                    break
                await asyncio.sleep(wait + random.random())
                continue

            ki, model = cands[0]
            self.rr = (ki + 1) % len(self.clients)
            try:
                async with self.sem:
                    resp = await asyncio.wait_for(
                        self.clients[ki].aio.models.generate_content(
                            model=model, contents=contents, config=cfg),
                        LLM_CALL_TIMEOUT)
                items = [schema(**d) for d in self._parse(resp.text)]
                return items
            except asyncio.CancelledError:
                raise
            except (json.JSONDecodeError, ValueError, TypeError) as e:   # bad/truncated JSON
                parse_fails += 1
                last_err = e
                log.warning("parse failure (%s, key %d): %s", model, ki, e)
                if parse_fails >= 3:
                    break
            except asyncio.TimeoutError as e:
                last_err = e
                self._cool(ki, model, 5)
            except Exception as e:  # google.genai.errors.APIError etc.
                last_err = e
                code = getattr(e, "code", None)
                msg = str(e)
                low = msg.lower()
                if code == 429 or "resource_exhausted" in low:
                    per_day = "perday" in low.replace(" ", "").replace("_", "") or "per day" in low
                    self._cool(ki, model, 1800 if per_day else 65)
                    log.warning("quota hit key=%d model=%s (per_day=%s)", ki, model, per_day)
                elif code == 404 or "not_found" in low:
                    log.error("model %s not found – removing it from the chain", model)
                    self.dead_models.add(model)
                elif code in (401, 403) or "api key" in low:
                    log.error("key %d rejected: %s", ki, msg[:200])
                    for m in self.models:
                        self._cool(ki, m, 3600)
                elif code in (500, 502, 503, 504) or "unavailable" in low or "overloaded" in low:
                    self._cool(None, model, 8 + random.random() * 6)   # whole model is busy
                    log.warning("model %s busy (%s)", model, code)
                else:
                    log.error("non-retryable LLM error (%s): %s", code, msg[:300])
                    raise
        raise Busy(str(last_err) if last_err else "all models busy")


POOL: LLMPool | None = None


# ───────────────────────── prompts ─────────────────────────
BASE_RULES = """You are an expert medical professor writing high-quality exam questions from lecture material.

RULES
- Read EVERYTHING provided (text and every attached slide image: figures, diagrams, tables, graphs, histology/radiology pictures). Do not skip sections.
- Test only the real scientific content. Ignore titles, objectives, agenda, references, course/university/doctor names, dates and any administrative text.
- BE SELECTIVE: do NOT make a question from every line. Ask about the main concepts and high-yield facts the lecture is built around (definitions, causes, mechanisms, key features, diagnosis, management, classifications, important numbers/cut-offs). Roughly 75% of the questions should cover those core points. Use the remaining ~25% for secondary details that are still exam-worthy (important exceptions, side effects, associations). Skip trivial, repeated or purely decorative details, long lists of examples, and one-off numbers with no clinical meaning.
- Use only facts supported by the material. Never invent facts.
- Write in English, keep standard medical terms.
- Give "correct_answer" (the single correct option) and exactly 3 "wrong_answers". Wrong answers must be plausible, from the same category and similar length as the correct one, and clearly wrong according to the material. Never use "all/none of the above" or "both A and B", and no A)/B) labels. Write "explanation" to justify exactly the text of correct_answer.
- Make sure correct_answer is truly correct according to the material and that the explanation supports exactly that text.
- question <= 280 characters, each option <= 90 characters, explanation <= 180 characters (say WHY it is correct; do not mention option letters).
- The student will NOT see the slides or any image. Use images/diagrams only as a source of facts, and write every question so it is fully answerable as plain text. Never write "in the image", "the figure", "the slide" or "as shown". If a finding comes from an image, describe it in words inside the question.
- Vary the type: recall, mechanism, differentiation, application. No duplicate questions.
"""

MODE_RULES = {
    "mcq": 'STYLE: direct MCQs. The "case" field must be "".',
    "cases": (
        'STYLE: every question is a CLINICAL CASE. Put a realistic vignette in "case" (2-5 sentences: age, sex, '
        "presentation, key examination/lab/imaging findings) and put the actual question in \"question\" "
        "(most likely diagnosis / mechanism / next step / expected finding / best drug...). The case must test "
        "the lecture content and must not reveal the answer. If the content is basic science, wrap it in a "
        "realistic clinical scenario."
    ),
    "auto": (
        'STYLE: decide from the content. If it is clinical (diseases, signs, investigations, management, pathology, '
        "pharmacology, microbiology, ...) make at least 60% of the questions clinical cases: vignette (2-5 sentences) in "
        '"case", the actual question in "question". For pure definitions, anatomy, histology, embryology or '
        'biochemistry facts use direct MCQs with "case" = "".'
    ),
}


@dataclass
class Page:
    no: int
    text: str = ""
    images: list = field(default_factory=list)   # [(bytes, mime)]


@dataclass
class Chunk:
    pages: list
    text: str
    images: dict   # slide_no -> (bytes, mime)  (first image of that slide, used for sending to user)
    parts: list    # interleaved str / Part for Gemini
    target: int


def build_chunks(pages: list[Page]) -> list[Chunk]:
    groups, cur, chars, imgs = [], [], 0, 0
    for p in pages:
        n_img = len(p.images)
        if cur and (len(cur) >= MAX_PAGES_PER_CHUNK or chars + len(p.text) > MAX_CHARS_PER_CHUNK
                    or imgs + n_img > MAX_IMAGES_PER_CHUNK):
            groups.append(cur)
            cur, chars, imgs = [], 0, 0
        cur.append(p)
        chars += len(p.text)
        imgs += n_img
    if cur:
        groups.append(cur)

    chunks = []
    for g in groups:
        text_total = sum(len(p.text) for p in g)
        n_img = sum(len(p.images) for p in g)
        if text_total < 150 and n_img == 0:
            continue  # nothing to ask about (title/blank slides)
        parts, first_imgs = [], {}
        for p in g:
            if p.text.strip():
                parts.append(f"=== SLIDE/PAGE {p.no} (text) ===\n{p.text.strip()}")
            for b, mime in p.images:
                parts.append(f"=== SLIDE/PAGE {p.no} (image) ===")
                parts.append(types.Part.from_bytes(data=b, mime_type=mime))
                first_imgs.setdefault(p.no, (b, mime))
        target = max(3, min(10, round(text_total / 550) + n_img))
        chunks.append(Chunk(g, "", first_imgs, parts, target))
    return chunks


# ───────────────────────── file extraction (CPU-bound → run in a thread) ─────────────────────────
def pages_from_pdf(data: bytes) -> tuple[list[Page], bool]:
    doc = fitz.open(stream=data, filetype="pdf")
    total = len(doc)
    pages = []
    for i in range(min(total, MAX_PAGES)):
        page = doc[i]
        text = page.get_text().strip()
        rect = page.rect
        area = max(rect.width * rect.height, 1)
        img_area = 0.0
        for info in page.get_image_info():          # real images only, measured by on-page size
            r = fitz.Rect(info["bbox"])
            r.intersect(rect)
            if not r.is_empty:
                img_area += r.width * r.height
        frac = img_area / area
        drawings = len(page.get_drawings()) if frac < 0.06 else 0   # vector diagrams / charts
        visual = frac > 0.06 or drawings > 25
        images = []
        if visual:
            zoom = min(2.0, 1500 / max(rect.width, rect.height))
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            images.append((pix.tobytes("jpeg", jpg_quality=75), "image/jpeg"))
        pages.append(Page(i + 1, text, images))
    doc.close()
    return pages, total > MAX_PAGES


def pages_from_pptx(data: bytes) -> tuple[list[Page], bool]:
    import io
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    pages = []
    for i, slide in enumerate(prs.slides, 1):
        if i > MAX_PAGES:
            break
        texts, pics = [], []
        for sh in walk(slide.shapes):
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = sh.image
                    if img.content_type in ("image/png", "image/jpeg") and len(img.blob) > 15_000:
                        pics.append((img.blob, img.content_type))   # >15KB: skips logos/icons
                except Exception:
                    pass
        pics.sort(key=lambda x: len(x[0]), reverse=True)
        pages.append(Page(i, "\n".join(texts), pics[:2]))
    return pages, len(prs.slides) > MAX_PAGES


def pages_from_text(text: str) -> list[Page]:
    paras = re.split(r"\n\s*\n", text)
    pages, buf, n = [], "", 1
    for p in paras:
        if buf and len(buf) + len(p) > 6000:
            pages.append(Page(n, buf))
            buf, n = "", n + 1
        buf += p + "\n\n"
    if buf.strip():
        pages.append(Page(n, buf))
    return pages


# ───────────────────────── telegram helpers ─────────────────────────
async def tg(fn, *a, **kw):
    for attempt in range(5):
        try:
            return await fn(*a, **kw)
        except BadRequest:
            raise
        except RetryAfter as e:
            await asyncio.sleep(float(e.retry_after) + 1)
        except (TimedOut, NetworkError):
            await asyncio.sleep(1.5 * (attempt + 1))
    return await fn(*a, **kw)


async def edit_status(msg, text: str):
    try:
        await msg.edit_text(text)
    except Exception:
        pass


# ───────────────────────── jobs ─────────────────────────
@dataclass
class Job:
    user_id: int
    chat_id: int
    mode: str
    kind: str          # pdf | pptx | image | txt
    file_id: str
    mime: str = ""


JOBQ: asyncio.Queue = asyncio.Queue()
ACTIVE: set[int] = set()
FINISH_LOCK = asyncio.Lock()


async def maybe_finish(bot, quiz_id: str):
    q = db.execute("SELECT * FROM quizzes WHERE quiz_id=?", (quiz_id,)).fetchone()
    if not q or q["finished"] or not q["gen_done"]:
        return
    r = db.execute("SELECT COUNT(*) t, COALESCE(SUM(answered),0) a, COALESCE(SUM(right),0) c "
                   "FROM polls WHERE quiz_id=?", (quiz_id,)).fetchone()
    if r["t"] == 0 or r["a"] < r["t"]:
        return
    db.execute("UPDATE quizzes SET finished=1 WHERE quiz_id=?", (quiz_id,))
    db.commit()
    total, correct = r["t"], r["c"]
    wrong = total - correct
    pct = round(correct / total * 100, 1)
    emoji = "🏆" if pct >= 80 else "📚" if pct >= 50 else "💪"
    await tg(bot.send_message, chat_id=q["chat_id"], text=(
        f"{emoji} Quiz Completed!\n\n✅ Correct: {correct}/{total}\n❌ Wrong:   {wrong}/{total}\n"
        f"📈 Score:   {pct}%\n\nSend another file to start a new quiz!"))


async def send_questions(bot, job: Job, quiz_id: str, qs: list[dict],
                         seen: set, counter: list) -> int:
    sent = 0
    for q in qs:
        key = re.sub(r"\W+", " ", q["question"].lower()).strip()
        if key in seen:
            continue
        seen.add(key)

        if q["case"]:
            await tg(bot.send_message, chat_id=job.chat_id, text="🏥 Case:\n" + q["case"][:4000])

        counter[0] += 1
        poll = await tg(
            bot.send_poll, chat_id=job.chat_id,
            question=clip(f"Q{counter[0]}: {q['question']}", 300),
            options=q["options"], type="quiz", correct_option_id=q["correct"],
            explanation=q["explanation"] or None, is_anonymous=True,   # hides who answered
        )
        db.execute("INSERT INTO polls(poll_id, quiz_id, user_id, chat_id, correct_idx) VALUES(?,?,?,?,?)",
                   (poll.poll.id, quiz_id, job.user_id, job.chat_id, q["correct"]))
        db.commit()
        sent += 1
        await asyncio.sleep(0.4)   # stay under Telegram's per-chat flood limit
    return sent


VERIFY = os.getenv("VERIFY", "0") == "1"   # optional extra answer check (doubles API calls) – off by default


def _norm(t: str) -> str:
    return re.sub(r"\W+", " ", t.lower()).strip()


async def verify_questions(chunk: Chunk, qs: list[dict]) -> tuple[list[dict], int, int]:
    """Blind second pass: another call answers every question WITHOUT seeing our answer key.
    Questions whose answer differs from the key are dropped. Returns (kept, dropped, unverified)."""
    if not qs:
        return qs, 0, 0
    lines = []
    for i, q in enumerate(qs):
        body = (q["case"] + "\n" if q["case"] else "") + q["question"]
        lines.append(f"[{i}] {body}\n" + "\n".join(f"   - {o}" for o in q["options"]))
    head = ("You are a strict medical examiner. Using ONLY the material below (text and images) plus standard "
            "medical knowledge, pick the single best answer for each numbered question. Return a JSON array "
            "with one item per question: id = the question number, answer = the EXACT text of the option you chose.\n\n"
            "QUESTIONS:\n" + "\n\n".join(lines) + "\n\nMATERIAL:")
    try:
        verdicts = await POOL.generate([head] + chunk.parts, schema=Verdict, temperature=0.0)
    except (Busy, Exception):
        log.warning("verification unavailable – sending unverified", exc_info=True)
        return qs, 0, len(qs)
    picked = {v.id: _norm(v.answer) for v in verdicts}
    kept, dropped, unverified = [], 0, 0
    for i, q in enumerate(qs):
        key = _norm(q["options"][q["correct"]])
        known = [_norm(o) for o in q["options"]]
        p = picked.get(i)
        if p is not None and p not in known:      # tolerate small wording differences
            close = [o for o in known if p and (o.startswith(p) or p.startswith(o) or p in o)]
            p = close[0] if len(close) == 1 else None
        if p is None:
            kept.append(q); unverified += 1       # no usable verdict → keep but report
        elif p == key:
            kept.append(q)
        else:
            dropped += 1                          # blind answer differs from the key → doubtful
    return kept, dropped, unverified


async def make_questions(job: "Job", chunk: Chunk) -> tuple[list[dict], int, int]:
    items = await POOL.generate(build_prompt(job, chunk))
    qs = [q for q in (normalize_question(it) for it in items) if q]
    if VERIFY:
        return await verify_questions(chunk, qs)
    return qs, 0, 0


def build_prompt(job: Job, chunk: Chunk) -> list:
    n = chunk.target
    head = (BASE_RULES + "\n" + MODE_RULES[job.mode] +
            f"\n\nGenerate about {n} questions for this part (at most {n + 2}); fewer is fine if the content is thin. "
            "Return a JSON array following the schema.\n\nMATERIAL:")
    return [head] + chunk.parts


async def run_job(app: Application, job: Job):
    bot = app.bot
    status = await tg(bot.send_message, chat_id=job.chat_id, text="⏳ Reading your file...")

    try:
        f = await tg(bot.get_file, job.file_id)
        data = bytes(await f.download_as_bytearray())
    except BadRequest:
        await edit_status(status, "❌ File too big (Telegram bots can read up to 20 MB). Split it or compress it.")
        return

    truncated = False
    try:
        if job.kind == "pdf":
            pages, truncated = await asyncio.to_thread(pages_from_pdf, data)
        elif job.kind == "pptx":
            pages, truncated = await asyncio.to_thread(pages_from_pptx, data)
        elif job.kind == "image":
            pages = [Page(1, "", [(data, job.mime or "image/jpeg")])]
        else:
            pages = pages_from_text(data.decode("utf-8", errors="ignore"))
    except Exception:
        log.exception("extract failed")
        await edit_status(status, "❌ Could not read this file. Is it corrupted or password-protected?")
        return

    chunks = build_chunks(pages)
    if not chunks:
        await edit_status(status, "❌ I couldn't find readable content in this file.")
        return

    quiz_id = uuid.uuid4().hex
    db.execute("INSERT INTO quizzes(quiz_id, user_id, chat_id, created) VALUES(?,?,?,?)",
               (quiz_id, job.user_id, job.chat_id, time.time()))
    db.commit()

    nchunks = len(chunks)
    await edit_status(status, f"🧠 Generating questions ({MODE_HELP[job.mode].split(' –')[0]}) – "
                              f"part 0/{nchunks} done. Questions start arriving shortly...")
    tasks = [asyncio.create_task(make_questions(job, c)) for c in chunks]
    dropped_total = unverified_total = 0
    seen: set = set()
    counter = [0]
    failed = 0
    try:
        for i, (chunk, task) in enumerate(zip(chunks, tasks), 1):
            try:
                qs, dropped, unverified = await task
                dropped_total += dropped
                unverified_total += unverified
            except Busy:
                failed += 1
                continue
            except Exception:
                log.exception("chunk failed")
                failed += 1
                continue
            await send_questions(bot, job, quiz_id, qs, seen, counter)
            await edit_status(status, f"🧠 Generating… part {i}/{nchunks} done ({counter[0]} questions so far)")
    finally:
        for t in tasks:
            t.cancel()

    async with FINISH_LOCK:
        db.execute("UPDATE quizzes SET gen_done=1 WHERE quiz_id=?", (quiz_id,))
        db.commit()
        await maybe_finish(bot, quiz_id)

    if counter[0] == 0:
        await edit_status(status, "⚠️ The AI servers are overloaded right now, no questions were generated. "
                                  "Please try again in a minute.")
        return
    note = f"✅ Done – {counter[0]} questions."
    if failed:
        note += f"\n⚠️ {failed} part(s) of the file were skipped because the AI was busy. Send the file again to cover them."
    if dropped_total:
        note += f"\n🔎 {dropped_total} doubtful question(s) were removed by the answer check."
    if unverified_total:
        note += f"\n⚠️ {unverified_total} question(s) could not be double-checked (AI busy) – review them."
    if truncated:
        note += f"\nℹ️ Only the first {MAX_PAGES} pages were used."
    await edit_status(status, note)


async def job_worker(app: Application):
    while True:
        job: Job = await JOBQ.get()
        try:
            await run_job(app, job)
        except asyncio.CancelledError:
            raise
        except Exception:
            log.exception("job crashed")
            try:
                await tg(app.bot.send_message, chat_id=job.chat_id,
                         text="❌ Something went wrong while processing your file. Please try again.")
            except Exception:
                pass
        finally:
            ACTIVE.discard(job.user_id)
            JOBQ.task_done()


# ───────────────────────── handlers ─────────────────────────
def keyboard():
    return ReplyKeyboardMarkup([list(WAITING_BUTTONS)], resize_keyboard=True, is_persistent=True)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    mode = get_mode(update.effective_user.id)
    await update.message.reply_text(
        "👋 Welcome to the Medical MCQ Bot!\n\n"
        "Just send me your lecture (PDF, PPTX, image or .txt) and I'll turn it into quiz questions – "
        "I also read the figures inside the slides.\n\n"
        "Pick the style with the buttons below:\n" + "\n".join(MODE_HELP.values()) +
        f"\n\nCurrent style: {MODE_HELP[mode].split(' –')[0]}",
        reply_markup=keyboard())


async def text_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    if text in WAITING_BUTTONS:
        mode = WAITING_BUTTONS[text]
        set_mode(update.effective_user.id, mode)
        await update.message.reply_text(f"✅ Style set: {MODE_HELP[mode]}\n\n📂 Now send your file (PDF / PPTX / image / txt).",
                                        reply_markup=keyboard())
    else:
        await update.message.reply_text("📂 Send me a file (PDF / PPTX / image / txt) to generate questions.",
                                        reply_markup=keyboard())


async def file_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    user_id, chat_id = update.effective_user.id, msg.chat_id
    doc = msg.document

    if msg.photo:
        kind, file_id, mime, size = "image", msg.photo[-1].file_id, "image/jpeg", msg.photo[-1].file_size or 0
    elif doc:
        name = (doc.file_name or "").lower()
        mt = doc.mime_type or ""
        file_id, size, mime = doc.file_id, doc.file_size or 0, mt
        if name.endswith(".pdf") or mt == "application/pdf":
            kind = "pdf"
        elif name.endswith(".pptx"):
            kind = "pptx"
        elif mt.startswith("image/"):
            kind = "image"
        elif name.endswith(".txt") or mt == "text/plain":
            kind = "txt"
        else:
            await msg.reply_text("❌ Unsupported file. Please send PDF, PPTX, an image or .txt")
            return
    else:
        return

    if size > TG_FILE_LIMIT:
        await msg.reply_text("❌ File is bigger than 20 MB (Telegram limit for bots). Split it or compress it.")
        return
    if user_id in ACTIVE:
        await msg.reply_text("⏳ I'm still working on your previous file. Wait for it to finish "
                             "(tip: send several pages as ONE PDF, not many photos).")
        return
    if JOBQ.qsize() >= MAX_QUEUE:
        await msg.reply_text("🚦 The bot is very busy right now. Please try again in a few minutes.")
        return

    ACTIVE.add(user_id)
    pos = JOBQ.qsize() + 1
    await JOBQ.put(Job(user_id, chat_id, get_mode(user_id), kind, file_id, mime))
    if pos > JOB_WORKERS:
        await msg.reply_text(f"🕐 You're in the queue (position {pos - JOB_WORKERS}). I'll start automatically.")


async def poll_answer_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    ans = update.poll_answer
    if not ans.option_ids:      # vote retracted
        return
    async with FINISH_LOCK:
        row = db.execute("SELECT * FROM polls WHERE poll_id=?", (ans.poll_id,)).fetchone()
        if not row or row["answered"] or row["user_id"] != ans.user.id:
            return
        right = int(ans.option_ids[0] == row["correct_idx"])
        db.execute("UPDATE polls SET answered=1, right=? WHERE poll_id=?", (right, ans.poll_id))
        db.commit()
        await maybe_finish(context.bot, row["quiz_id"])


async def poll_state_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Anonymous polls send no PollAnswer, but Telegram sends the poll state when someone votes.
    The quiz is private (one student per chat), so the option that got a vote is the student's answer."""
    poll = update.poll
    if poll is None or poll.total_voter_count < 1:
        return
    async with FINISH_LOCK:
        row = db.execute("SELECT * FROM polls WHERE poll_id=?", (poll.id,)).fetchone()
        if not row or row["answered"] or row["chat_id"] < 0:   # groups: many voters, no personal score
            return
        chosen = max(range(len(poll.options)), key=lambda i: poll.options[i].voter_count)
        if poll.options[chosen].voter_count < 1:
            return
        right = int(chosen == row["correct_idx"])
        db.execute("UPDATE polls SET answered=1, right=? WHERE poll_id=?", (right, poll.id))
        db.commit()
        await maybe_finish(context.bot, row["quiz_id"])


async def error_handler(update, context: ContextTypes.DEFAULT_TYPE):
    log.error("handler error: %s", context.error, exc_info=context.error)


# ───────────────────────── lifecycle ─────────────────────────
async def health_server(port: int):
    """Tiny HTTP responder for hosts that insist on an open port / uptime pings."""
    async def handle(reader, writer):
        try:
            await reader.read(1024)
            writer.write(b"HTTP/1.1 200 OK\r\nContent-Type: text/plain\r\nContent-Length: 2\r\nConnection: close\r\n\r\nok")
            await writer.drain()
        finally:
            writer.close()
    return await asyncio.start_server(handle, "0.0.0.0", port)


async def post_init(app: Application):
    global POOL
    POOL = LLMPool(GEMINI_KEYS, GEMINI_MODELS, LLM_CONCURRENCY)
    app.bot_data["workers"] = [asyncio.create_task(job_worker(app)) for _ in range(JOB_WORKERS)]
    port = os.getenv("PORT")
    if port and os.getenv("HEALTH_SERVER", "1") == "1":
        app.bot_data["health"] = await health_server(int(port))
    log.info("Bot ready: %d key(s), models=%s, workers=%d", len(GEMINI_KEYS), GEMINI_MODELS, JOB_WORKERS)


def main():
    logging.basicConfig(format="%(asctime)s %(levelname)s %(name)s: %(message)s", level=logging.INFO)
    for noisy in ("httpx", "httpcore", "telegram", "google_genai", "hpack"):
        logging.getLogger(noisy).setLevel(logging.WARNING)   # httpx would log the bot token in URLs
    if not TELEGRAM_TOKEN:
        raise SystemExit("TELEGRAM_TOKEN is not set")

    app = (Application.builder().token(TELEGRAM_TOKEN)
           .concurrent_updates(True)          # one slow user no longer blocks the others
           .post_init(post_init).build())
    app.add_handler(CommandHandler(["start", "help", "mode"], start))
    app.add_handler(PollAnswerHandler(poll_answer_handler))
    app.add_handler(PollHandler(poll_state_handler))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, text_handler))
    app.add_handler(MessageHandler(filters.Document.ALL | filters.PHOTO, file_handler))
    app.add_error_handler(error_handler)
    app.run_polling(allowed_updates=[Update.MESSAGE, Update.POLL_ANSWER, Update.POLL])


if __name__ == "__main__":
    main()
